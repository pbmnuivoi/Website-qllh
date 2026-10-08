# -*- coding: utf-8 -*-
"""Fallback server for SCHOOL APP - CSDL EXCEL TONG.
Uses only Python standard library + openpyxl. It mirrors the Node API used by the frontend.
"""
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from urllib.parse import urlparse
from pathlib import Path
import json, os, re, secrets, threading, traceback
from http import HTTPStatus

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("THIEU THU VIEN openpyxl. Hay cai Python va chay: py -m pip install openpyxl")
    raise

ROOT = Path(__file__).resolve().parent
PUBLIC_DIR = ROOT / 'school-app'
DB_DIR = ROOT / 'db'
DB_FILE = DB_DIR / 'CSDL_EXCEL_TONG.xlsx'
PORT = int(os.environ.get('PORT', '3000'))

DEFAULT_USERS = [{'username':'Admin','password':'@Zz123456','role':'admin','name':'Quản trị viên'}]
DEFAULT_CLASSES = ['1A','1B','1C','1D','1E','2A','2B','2C','2D','2E','3A','3B','3C','3D','3E','4A','4B','4C','4D','4E','5A','5B','5C','5D','5E']
DEFAULT_SUBJECTS = ['Tin học','Công nghệ']
DEFAULT_SETTINGS = {'schoolName':'Trường Tiểu học Núi Voi','ubndName':'UBND phường Chi Lăng'}

db_lock = threading.RLock()
sessions = {}
db = {'users': [], 'classes': [], 'subjects': [], 'students': {}, 'scores': {}, 'settings': DEFAULT_SETTINGS.copy()}


def clone(v):
    return json.loads(json.dumps(v, ensure_ascii=False))


def clean_student_name(name, sid=''):
    if not name:
        return 'Chưa có tên'
    n = str(name).strip()
    if sid and n.startswith(sid):
        n = n[len(sid):].strip()
    n = re.sub(r'^\d{8,}\s+', '', n).strip()
    return n or 'Chưa có tên'


def normalize_students(items):
    out=[]
    for i, st in enumerate(items if isinstance(items,list) else []):
        sid=str(st.get('id', st.get('ma',''))).strip()
        if sid:
            out.append({'id':sid,'name':clean_student_name(st.get('name',st.get('hoTen','')),sid),'stt':st.get('stt') or i+1})
    return out


def sort_classes(items):
    def key(x):
        m=re.match(r'^(\d+)(.*)$',str(x))
        return (int(m.group(1)) if m else 999, str(x))
    return sorted(items, key=key)


def ensure_score_shape():
    db.setdefault('scores', {})
    subjects=db.get('subjects') or DEFAULT_SUBJECTS
    for students in db.get('students',{}).values():
        for st in students:
            sid=st['id']
            if sid not in db['scores'] or not isinstance(db['scores'][sid],dict):
                db['scores'][sid]={}
            entry=db['scores'][sid]
            if isinstance(entry.get('points'), (int,float)):
                old_points=entry.get('points',10); old_history=entry.get('history',[])
                db['scores'][sid]={}
                for j,sub in enumerate(subjects):
                    db['scores'][sid][sub]={'points':old_points if j==0 else 10,'history':old_history if j==0 else []}
            for sub in subjects:
                e=entry.get(sub) if isinstance(entry,dict) else None
                if not isinstance(e,dict) or not isinstance(e.get('points'),(int,float)):
                    db['scores'][sid][sub]={'points':10,'history':[]}


def normalize_db(inp):
    global db
    d=inp or {}
    out={
        'users': d.get('users') if isinstance(d.get('users'),list) and d.get('users') else clone(DEFAULT_USERS),
        'classes': d.get('classes') if isinstance(d.get('classes'),list) and d.get('classes') else clone(DEFAULT_CLASSES),
        'subjects': [str(x).strip() for x in (d.get('subjects') if isinstance(d.get('subjects'),list) and d.get('subjects') else DEFAULT_SUBJECTS) if str(x).strip()],
        'students': d.get('students') if isinstance(d.get('students'),dict) else {},
        'scores': d.get('scores') if isinstance(d.get('scores'),dict) else {},
        'settings': {**clone(DEFAULT_SETTINGS), **(d.get('settings') if isinstance(d.get('settings'),dict) else {})}
    }
    out['classes']=sort_classes(list(dict.fromkeys(str(x).strip() for x in out['classes'] if str(x).strip())))
    for cls in list(out['students']): out['students'][cls]=normalize_students(out['students'][cls])
    old_db=db
    db=out
    ensure_score_shape()
    out=db
    db=old_db
    return out


def rows(ws):
    if ws is None or ws.max_row < 2: return []
    headers=[c.value for c in ws[1]]
    out=[]
    for row in ws.iter_rows(min_row=2, values_only=True):
        out.append({str(headers[i]): row[i] if i < len(row) else '' for i in range(len(headers))})
    return out


def read_workbook():
    if not DB_FILE.exists(): return None
    wb=load_workbook(DB_FILE, read_only=True, data_only=False)
    out={'users':[],'classes':[],'subjects':[],'students':{},'scores':{},'settings':clone(DEFAULT_SETTINGS)}
    try:
        for r in rows(wb['USERS'] if 'USERS' in wb.sheetnames else None):
            out['users'].append({'username':str(r.get('username') or ''),'password':str(r.get('password') or ''),'role':str(r.get('role') or 'admin'),'name':str(r.get('name') or '')})
        for r in rows(wb['CLASSES'] if 'CLASSES' in wb.sheetnames else None):
            if r.get('classCode'): out['classes'].append(str(r['classCode']))
        for r in rows(wb['SUBJECTS'] if 'SUBJECTS' in wb.sheetnames else None):
            if r.get('subject'): out['subjects'].append(str(r['subject']))
        for r in rows(wb['STUDENTS'] if 'STUDENTS' in wb.sheetnames else None):
            cls=str(r.get('classCode') or '').strip()
            if not cls: continue
            out['students'].setdefault(cls,[]).append({'id':str(r.get('id') or '').strip(),'name':str(r.get('name') or '').strip(),'stt':r.get('stt') or len(out['students'][cls])+1})
        if 'SETTINGS' in wb.sheetnames:
            for r in rows(wb['SETTINGS']):
                k=str(r.get('key') or '').strip(); v=str(r.get('value') or '').strip()
                if k: out.setdefault('settings',{})[k]=v
        for r in rows(wb['SCORES'] if 'SCORES' in wb.sheetnames else None):
            sid=str(r.get('studentId') or '').strip(); sub=str(r.get('subject') or '').strip()
            if not sid or not sub: continue
            try: hist=json.loads(str(r.get('history') or '[]'))
            except Exception: hist=[]
            out['scores'].setdefault(sid,{})[sub]={'points':float(r.get('points') if r.get('points') is not None else 10),'history':hist if isinstance(hist,list) else []}
    finally:
        wb.close()
    return normalize_db(out)


def write_workbook(next_db):
    global db
    with db_lock:
        data=clone(next_db)
        normalize_db(data)
        wb=Workbook()
        default=wb.active
        wb.remove(default)
        sheets={
            'USERS': [('username','password','role','name')],
            'CLASSES': [('classCode',)],
            'SUBJECTS': [('subject',)],
            'STUDENTS': [('classCode','stt','id','name')],
            'SCORES': [('studentId','subject','points','history')],
            'SETTINGS': [('key','value')]
        }
        for name, headers in sheets.items():
            ws=wb.create_sheet(name); ws.append(headers[0])
        for u in data['users']: wb['USERS'].append([u.get('username',''),u.get('password',''),u.get('role','admin'),u.get('name','')])
        for c in data['classes']: wb['CLASSES'].append([c])
        for s in data['subjects']: wb['SUBJECTS'].append([s])
        for c in sort_classes(data['students']):
            for st in data['students'][c]: wb['STUDENTS'].append([c,st.get('stt',1),st.get('id',''),st.get('name','')])
        for k,v in data.get('settings',{}).items(): wb['SETTINGS'].append([k,v])
        for sid,item in data['scores'].items():
            for sub,e in item.items():
                if sub == '_classCode': continue
                wb['SCORES'].append([sid,sub,float(e.get('points',10)),json.dumps(e.get('history',[]) if isinstance(e.get('history',[]),list) else [],ensure_ascii=False)])
        DB_DIR.mkdir(parents=True, exist_ok=True)
        tmp=DB_FILE.with_name(DB_FILE.stem + '.tmp.xlsx')
        wb.save(tmp)
        if DB_FILE.exists(): DB_FILE.unlink()
        tmp.replace(DB_FILE)
        db=normalize_db(data)


def seed_from_json_files():
    data_dir=PUBLIC_DIR/'data'
    if data_dir.exists():
        for f in data_dir.glob('students_*.json'):
            try:
                x=json.loads(f.read_text(encoding='utf-8'))
                if x.get('classCode') and isinstance(x.get('students'),list):
                    db['students'][x['classCode']]=normalize_students(x['students'])
                    if x['classCode'] not in db['classes']: db['classes'].append(x['classCode'])
            except Exception as e: print('Không đọc được',f,e)
    db['classes']=sort_classes(db['classes']); ensure_score_shape()


def initialize():
    global db
    existing=read_workbook()
    if existing and existing.get('users'):
        db=existing
    else:
        db={'users':clone(DEFAULT_USERS),'classes':clone(DEFAULT_CLASSES),'subjects':clone(DEFAULT_SUBJECTS),'students':{},'scores':{},'settings':clone(DEFAULT_SETTINGS)}
        seed_from_json_files(); write_workbook(db)


def cookie_sid(handler):
    raw=handler.headers.get('Cookie','')
    for part in raw.split(';'):
        part=part.strip()
        if part.startswith('school_sid='): return part.split('=',1)[1]
    return None


def get_session(handler):
    sid=(handler.headers.get('X-Session-Id') or cookie_sid(handler) or '').strip()
    return sessions.get(sid)


def send_json(handler, code, obj, extra=None):
    body=json.dumps(obj,ensure_ascii=False).encode('utf-8')
    handler.send_response(code); handler.send_header('Content-Type','application/json; charset=utf-8'); handler.send_header('Content-Length',str(len(body)))
    handler.send_header('Cache-Control','no-store, no-cache, must-revalidate, proxy-revalidate'); handler.send_header('Pragma','no-cache'); handler.send_header('Expires','0')
    origin=handler.headers.get('Origin')
    if origin:
        handler.send_header('Access-Control-Allow-Origin',origin); handler.send_header('Access-Control-Allow-Credentials','true')
    if extra:
        for k,v in extra.items(): handler.send_header(k,v)
    handler.end_headers(); handler.wfile.write(body)


class Handler(SimpleHTTPRequestHandler):
    extensions_map={**SimpleHTTPRequestHandler.extensions_map,' .js':'application/javascript'}
    def translate_path(self,path):
        rel=urlparse(path).path.lstrip('/')
        if not rel: rel='index.html'
        return str((PUBLIC_DIR/rel).resolve())
    def end_headers(self):
        p=urlparse(self.path).path
        if p.startswith('/api/') or re.search(r'\.(html|js|css|json)$',p,re.I):
            self.send_header('Cache-Control','no-store, no-cache, must-revalidate, proxy-revalidate'); self.send_header('Pragma','no-cache'); self.send_header('Expires','0')
        super().end_headers()
    def do_OPTIONS(self):
        self.send_response(204); self.send_header('Access-Control-Allow-Origin',self.headers.get('Origin','null')); self.send_header('Access-Control-Allow-Credentials','true'); self.send_header('Access-Control-Allow-Headers','Content-Type, X-Session-Id'); self.send_header('Access-Control-Allow-Methods','GET,POST,OPTIONS'); self.end_headers()
    def read_body(self):
        n=int(self.headers.get('Content-Length','0')); return self.rfile.read(n) if n else b''
    def do_GET(self):
        p=urlparse(self.path).path
        try:
            if p=='/api/health': return send_json(self,200,{'ok':True,'database':'CSDL_EXCEL_TONG.xlsx','server':'python'})
            if p=='/api/session': return send_json(self,200,{'ok':True,'session':clone(get_session(self)) if get_session(self) else None})
            if p=='/api/bootstrap':
                latest=read_workbook()
                if latest:
                    with db_lock: globals()['db']=latest
                s=get_session(self); sid=self.headers.get('X-Session-Id') or cookie_sid(self)
                return send_json(self,200,{'ok':True,'data':clone(db),'session':clone(s) if s else None,'sessionToken':sid})
            if p=='/api/export':
                if not DB_FILE.exists(): return send_json(self,404,{'ok':False,'msg':'Chưa có CSDL Excel'})
                data=DB_FILE.read_bytes(); self.send_response(200); self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); self.send_header('Content-Disposition','attachment; filename="CSDL_EXCEL_TONG.xlsx"'); self.send_header('Content-Length',str(len(data))); self.end_headers(); self.wfile.write(data); return
            return super().do_GET()
        except Exception as e:
            traceback.print_exc(); return send_json(self,500,{'ok':False,'msg':str(e)})
    def do_POST(self):
        p=urlparse(self.path).path
        try:
            raw=self.read_body(); body=json.loads(raw.decode('utf-8') or '{}') if raw else {}
            if p=='/api/login':
                username=str(body.get('username','')).strip(); password=str(body.get('password',''))
                user=next((u for u in db['users'] if u.get('username')==username and u.get('password')==password),None)
                if not user: return send_json(self,401,{'ok':False,'msg':'Sai tài khoản hoặc mật khẩu'})
                sid=secrets.token_hex(24); session={'username':user['username'],'name':user.get('name') or user['username'],'role':user.get('role','admin'),'loginAt':int(__import__('time').time()*1000)}; sessions[sid]=session
                return send_json(self,200,{'ok':True,'user':clone(session),'sessionToken':sid},{'Set-Cookie':f'school_sid={sid}; HttpOnly; Path=/; SameSite=Lax'})
            if p=='/api/logout':
                sid=(self.headers.get('X-Session-Id') or cookie_sid(self) or '').strip(); sessions.pop(sid,None)
                return send_json(self,200,{'ok':True},{'Set-Cookie':'school_sid=; HttpOnly; Path=/; Max-Age=0; SameSite=Lax'})
            if p=='/api/save':
                s=get_session(self)
                if not s or s.get('role')!='admin': return send_json(self,401,{'ok':False,'msg':'Vui lòng đăng nhập Admin.'})
                write_workbook(body.get('data') or {}); return send_json(self,200,{'ok':True})
            return send_json(self,404,{'ok':False,'msg':'API không tồn tại'})
        except Exception as e:
            traceback.print_exc(); return send_json(self,500,{'ok':False,'msg':str(e)})


def main():
    initialize()
    print('='*55); print('TRUONG TIEU HOC - CSDL EXCEL TONG'); print('='*55)
    print(f'Server Python đang chạy tại http://localhost:{PORT}/')
    print(f'CSDL: {DB_FILE}')
    print('Health: http://localhost:%d/api/health' % PORT)
    ThreadingHTTPServer(('0.0.0.0',PORT),Handler).serve_forever()

if __name__=='__main__': main()
